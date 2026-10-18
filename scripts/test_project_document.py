import tempfile
from pathlib import Path
import unittest
from unittest.mock import patch

from docx import Document
from docx.shared import RGBColor

from build_project_document import configure, normalize_fonts, render_markdown, validate


class ProjectDocumentTests(unittest.TestCase):
    def test_vietnamese_headings_table_and_examples_retain_requested_font(self):
        document = Document()
        configure(document)
        render_markdown(document, "# Kiến trúc và bảo mật\n\nTiếng Việt, quyền và dữ liệu.\n\n| Cột | Nội dung |\n| --- | --- |\n| DB | Không chứa API key |\n\n```text\nreconx --help\n```\n")
        normalize_fonts(document)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "manual.docx"
            document.save(output)
            result = validate(output)
            self.assertEqual(result["xml_validation"], "PASS")
            self.assertGreater(result["text_characters"], 80)
            restored = Document(output)
            self.assertIn("Tiếng Việt", restored.paragraphs[1].text)
            self.assertEqual(restored.tables[0].cell(1, 1).text, "Không chứa API key")

    def test_validation_rejects_nonblack_text(self):
        document = Document()
        configure(document)
        run = document.add_paragraph().add_run("Must not be blue")
        from build_project_document import font
        font(run._r.get_or_add_rPr())
        run.font.color.rgb = RGBColor(0, 0, 255)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "bad.docx"
            document.save(output)
            with self.assertRaisesRegex(ValueError, "Nonblack color"):
                validate(output)

    def test_existing_document_is_never_replaced(self):
        from build_project_document import main
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "existing.docx"
            output.write_bytes(b"User document must survive")
            with patch("sys.argv", ["build_project_document.py", "--output", str(output)]):
                with self.assertRaises(SystemExit) as result:
                    main()
            self.assertEqual(result.exception.code, 2)
            self.assertEqual(output.read_bytes(), b"User document must survive")


if __name__ == "__main__":
    unittest.main()
